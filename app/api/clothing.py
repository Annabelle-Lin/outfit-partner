# -*- coding: utf-8 -*-
"""衣物模块（计划书 13.2）。

POST   /clothing           新增衣物
GET    /clothing           获取衣橱列表
GET    /clothing/{item_id} 获取单件详情
PUT    /clothing/{item_id} 修改衣物
DELETE /clothing/{item_id} 删除衣物
POST   /clothing/import    Excel 批量导入（增量三）
GET    /clothing/template  下载导入模板（增量三）
GET    /clothing/meta      获取标签枚举
"""
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app import schemas, utils
from app.agent.engine import rule_repo
from app.api import ok
from app.db import ClothingItem, User, get_db

router = APIRouter()


def _item_to_response(item: ClothingItem) -> dict:
    return {
        "item_id": item.item_id,
        "user_id": item.user_id,
        "name": item.name,
        "category": item.category,
        "sub_category": item.sub_category,
        "color": item.color,
        "color_brightness": item.color_brightness,
        "season": item.season,
        "fit_snug": item.fit_snug,
        "fit_silhouette": item.fit_silhouette,
        "fabric": item.fabric,
        "style": item.style,
        "occasion": item.occasion or [],
        "warmth_level": item.warmth_level,
        "formality_level": item.formality_level,
        "image_url": item.image_url,
        "tags": item.tags or [],
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
    }


@router.get("/meta", response_model=schemas.APIResponse)
def get_meta():
    """标签枚举。前端录入表单的下拉选项直接用这个接口。"""
    meta = rule_repo.clothing_meta
    if not meta:
        raise HTTPException(status_code=500, detail="规则库未加载")
    return ok(data=meta)


# Excel 导入字段（与 ClothingItem 模型一一对应，必填项见 REQUIRED_FIELDS）
IMPORT_FIELDS = [
    "name", "category", "sub_category", "color", "color_brightness",
    "season", "fit_snug", "fit_silhouette", "fabric", "style",
    "occasion", "warmth_level", "formality_level", "image_url", "tags",
]
REQUIRED_FIELDS = ["name", "category", "sub_category", "color"]


@router.get("/template")
def download_template():
    """下载 Excel 导入模板：首页为待填写表格，第二页为枚举填写说明。"""
    meta = rule_repo.clothing_meta or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "衣物导入"
    ws.append(IMPORT_FIELDS)
    # 示例行，导入时会跳过表头，用户删除示例即可
    ws.append(["白色棉T恤", "上装", "T恤", "白", "亮色", "四季", "合身", "直筒", "棉", "基础款",
               "通勤/日常", 2, 2, "", "百搭/基础"])

    ws2 = wb.create_sheet("填写说明")
    ws2.append(["字段", "可选值（与 /clothing/meta 一致）"])
    cat_map = meta.get("category", {})
    if cat_map:
        ws2.append(["category", " / ".join(cat_map.keys())])
        for cat, subs in cat_map.items():
            ws2.append([f"sub_category（{cat}）", " / ".join(str(s) for s in subs)])
    for key in ("color_family", "color_brightness", "season", "fit_snug",
                "fit_silhouette", "fabric", "style", "occasion"):
        vals = meta.get(key)
        if vals:
            ws2.append([key, " / ".join(str(v) for v in vals)])
    ws2.append(["说明", "name/category/sub_category/color 必填；occasion、tags 多个值用 / 分隔；"
                        "warmth_level、formality_level 填 1-5 的整数，留空默认 3"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=wardrobe_import_template.xlsx"},
    )


@router.post("/import")
async def import_items(user_id: str, file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Excel 批量导入：逐行校验，成功行直接入库，失败行返回行号与原因。"""
    user = db.query(User).filter(User.user_id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="请上传 .xlsx 格式文件（可先下载导入模板）")

    raw = await file.read()
    try:
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="文件解析失败，请使用标准模板")

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 1:
        raise HTTPException(status_code=400, detail="文件内容为空")

    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col_map = {f: header.index(f) for f in IMPORT_FIELDS if f in header}
    missing_cols = [f for f in REQUIRED_FIELDS if f not in col_map]
    if missing_cols:
        raise HTTPException(status_code=400, detail=f"模板缺少必填列：{', '.join(missing_cols)}")

    meta = rule_repo.clothing_meta or {}
    cat_map = meta.get("category", {})

    def cell(row, field):
        idx = col_map.get(field)
        if idx is None or idx >= len(row):
            return ""
        v = row[idx]
        return str(v).strip() if v is not None else ""

    def cell_list(row, field):
        text = cell(row, field)
        if not text:
            return []
        for sep in ("，", ",", "、", "；", ";"):
            text = text.replace(sep, "/")
        return [p.strip() for p in text.split("/") if p.strip()]

    def parse_level(row, field):
        """解析 1-5 整数等级，留空默认 3；返回 (值, 错误原因)。"""
        text = cell(row, field)
        if not text:
            return 3, None
        try:
            iv = int(float(text))
        except ValueError:
            return None, f"{field} 应为 1-5 的整数，实际为「{text}」"
        if not 1 <= iv <= 5:
            return None, f"{field} 超出范围（1-5），实际为 {iv}"
        return iv, None

    success = 0
    errors = []
    for row_no, row in enumerate(rows[1:], start=2):
        if all(c is None or str(c).strip() == "" for c in row):
            continue  # 跳过空行
        data = {f: cell(row, f) for f in IMPORT_FIELDS}

        missing_req = [f for f in REQUIRED_FIELDS if not data[f]]
        if missing_req:
            errors.append({"row": row_no, "reason": f"必填字段为空：{', '.join(missing_req)}"})
            continue
        if cat_map and data["category"] not in cat_map:
            errors.append({"row": row_no,
                           "reason": f"品类「{data['category']}」不存在，可选：{'/'.join(cat_map.keys())}"})
            continue
        subs = cat_map.get(data["category"], [])
        if subs and data["sub_category"] not in subs:
            errors.append({"row": row_no,
                           "reason": f"品类「{data['category']}」下没有二级品类「{data['sub_category']}」"})
            continue
        warmth, err = parse_level(row, "warmth_level")
        if err:
            errors.append({"row": row_no, "reason": err})
            continue
        formality, err = parse_level(row, "formality_level")
        if err:
            errors.append({"row": row_no, "reason": err})
            continue

        item = ClothingItem(
            item_id=utils.gen_id("i"),
            user_id=user_id,
            name=data["name"],
            category=data["category"],
            sub_category=data["sub_category"],
            color=data["color"],
            color_brightness=data["color_brightness"],
            season=data["season"] or "四季",
            fit_snug=data["fit_snug"],
            fit_silhouette=data["fit_silhouette"],
            fabric=data["fabric"],
            style=data["style"],
            occasion=cell_list(row, "occasion"),
            warmth_level=warmth,
            formality_level=formality,
            image_url=data["image_url"],
            tags=cell_list(row, "tags"),
        )
        db.add(item)
        success += 1

    db.commit()
    return ok(data={"success": success, "failed": len(errors), "errors": errors})


@router.post("", response_model=schemas.APIResponse)
def create_item(req: schemas.ClothingItemCreate, user_id: str, db: Session = Depends(get_db)):
    """新增衣物。"""
    user = db.query(User).filter(User.user_id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")

    item = ClothingItem(
        item_id=utils.gen_id("i"),
        user_id=user_id,
        name=req.name,
        category=req.category,
        sub_category=req.sub_category,
        color=req.color,
        color_brightness=req.color_brightness,
        season=req.season,
        fit_snug=req.fit_snug,
        fit_silhouette=req.fit_silhouette,
        fabric=req.fabric,
        style=req.style,
        occasion=req.occasion,
        warmth_level=req.warmth_level,
        formality_level=req.formality_level,
        image_url=req.image_url,
        tags=req.tags,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return ok(data=_item_to_response(item))


@router.get("", response_model=schemas.APIResponse)
def list_items(
    user_id: str,
    category: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
):
    """衣橱列表，支持按品类筛选。"""
    q = db.query(ClothingItem).filter(ClothingItem.user_id == user_id)
    if category:
        q = q.filter(ClothingItem.category == category)
    total = q.count()
    items = q.order_by(ClothingItem.created_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    return ok(data=schemas.ClothingListResponse(
        total=total,
        page=page,
        page_size=page_size,
        items=[_item_to_response(i) for i in items],
    ).model_dump())


@router.get("/{item_id}", response_model=schemas.APIResponse)
def get_item(item_id: str, db: Session = Depends(get_db)):
    item = db.query(ClothingItem).filter(ClothingItem.item_id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="衣物不存在")
    return ok(data=_item_to_response(item))


@router.put("/{item_id}", response_model=schemas.APIResponse)
def update_item(item_id: str, req: schemas.ClothingItemUpdate, db: Session = Depends(get_db)):
    item = db.query(ClothingItem).filter(ClothingItem.item_id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="衣物不存在")

    update_data = req.model_dump(exclude_unset=True)
    for k, v in update_data.items():
        setattr(item, k, v)
    db.commit()
    db.refresh(item)
    return ok(data=_item_to_response(item))


@router.delete("/{item_id}", response_model=schemas.APIResponse)
def delete_item(item_id: str, db: Session = Depends(get_db)):
    item = db.query(ClothingItem).filter(ClothingItem.item_id == item_id).first()
    if not item:
        raise HTTPException(status_code=404, detail="衣物不存在")
    db.delete(item)
    db.commit()
    return ok(data={"deleted": True})


